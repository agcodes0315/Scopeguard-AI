"""Strict local document extraction. No OCR or external upload service."""
from io import BytesIO
from pathlib import Path

MAX_BYTES = 3 * 1024 * 1024
MAX_CHARS = 30000
MAX_PAGES = 25


def extract_document(filename: str, content: bytes) -> dict:
    ext = Path(filename).suffix.lower()
    if ext not in {'.pdf', '.docx', '.txt'}:
        raise ValueError('Only PDF, DOCX, and TXT agreements are supported.')
    if not content or len(content) > MAX_BYTES:
        raise ValueError('Agreement must be nonempty and no larger than 3 MB.')
    if ext == '.pdf':
        if not content.startswith(b'%PDF-'):
            raise ValueError('File does not appear to be a valid PDF.')
        import fitz
        try:
            with fitz.open(stream=content, filetype='pdf') as doc:
                if len(doc) > MAX_PAGES: raise ValueError('PDF exceeds 25 pages.')
                if doc.needs_pass: raise ValueError('Password-protected PDF is not supported.')
                text = '\n'.join(f'[Page {i+1}]\n' + page.get_text(sort=True) for i, page in enumerate(doc))
                pages = len(doc)
        except (RuntimeError, TypeError) as exc:
            raise ValueError('PDF could not be parsed.') from exc
    elif ext == '.docx':
        if not content.startswith(b'PK'):
            raise ValueError('File does not appear to be a DOCX document.')
        from docx import Document
        try:
            doc = Document(BytesIO(content))
            sections = [x.text for x in doc.paragraphs]
            for table in doc.tables:
                for row in table.rows:
                    sections.append(' | '.join(cell.text for cell in row.cells))
            text = '\n'.join(sections)
        except Exception as exc:
            raise ValueError('DOCX could not be parsed.') from exc
        pages = None
    else:
        try: text = content.decode('utf-8-sig')
        except UnicodeError as exc: raise ValueError('TXT must use UTF-8 encoding.') from exc
        pages = None
    text = text.strip()
    if len(text) < 20:
        raise ValueError('Document contains too little selectable text. Scanned PDFs require OCR, not included in this MVP.')
    if len(text) > MAX_CHARS:
        raise ValueError('Extracted document exceeds 30,000 characters; upload a shorter agreement.')
    return {'filename': Path(filename).name, 'text': text, 'characters': len(text), 'pages': pages, 'stored': False}
